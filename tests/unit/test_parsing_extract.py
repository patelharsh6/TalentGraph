"""Steps 2.2–2.6: type detection, limits and the PDF/DOCX/TXT extractors (exit criteria 2-A, 2-E)."""

import io

import docx
import pymupdf
import pytest
from docx.oxml import parse_xml
from parsing_samples import SAMPLES, Sample
from rapidfuzz import fuzz

from talentgraph.agents.parsing.detect import check_size, detect_file_type
from talentgraph.agents.parsing.docx import extract_docx
from talentgraph.agents.parsing.pdf import extract_pdf
from talentgraph.agents.parsing.text import decode_text, extract_text
from talentgraph.core.errors import FileTooLarge, UnsupportedContent, UnsupportedFileType

EXPECTED_LAYOUT = {
    ("pdf", "single_column"): "single",
    ("pdf", "two_column"): "multi",
    ("pdf", "sidebar"): "multi",
    ("pdf", "table_based"): "table",
    ("pdf", "creative"): "creative",
    ("docx", "two_column"): "multi",
    ("docx", "sidebar"): "multi",
    ("docx", "single_column"): "single",
    ("docx", "creative"): "single",
    ("txt", "minimal_txt"): "single",
}


def test_fixtures_cover_every_template_and_format() -> None:
    templates = {s.template for s in SAMPLES}
    formats = {s.path.suffix for s in SAMPLES}
    assert templates == {
        "single_column",
        "two_column",
        "sidebar",
        "table_based",
        "creative",
        "minimal_txt",
    }
    assert formats == {".pdf", ".docx", ".txt"}


# -------------------------------------------------------------------------------------------------
# Detection and limits
# -------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda s: s.id)
def test_detects_type_from_content(sample: Sample) -> None:
    assert detect_file_type(sample.path.read_bytes()) == sample.path.suffix.lstrip(".")


@pytest.mark.parametrize(
    "data",
    [
        b"\x89PNG\r\n\x1a\n" + b"\x00" * 100,
        b"\xff\xd8\xff\xe0" + b"\x00" * 100,
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100,  # legacy .doc
        b"\x00\x01\x02\x03" * 100,
        b"   ",
    ],
    ids=["png", "jpeg", "doc", "binary", "empty"],
)
def test_rejects_unsupported_types(data: bytes) -> None:
    with pytest.raises(UnsupportedFileType) as exc:
        detect_file_type(data)
    assert exc.value.http_status == 415


def test_rejects_zip_that_is_not_docx() -> None:
    buf = io.BytesIO()
    import zipfile

    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(UnsupportedFileType):
        detect_file_type(buf.getvalue())


def test_size_limit_is_413() -> None:
    check_size(b"x" * (10 * 1024 * 1024), 10)
    with pytest.raises(FileTooLarge) as exc:
        check_size(b"%PDF" + b"x" * (11 * 1024 * 1024), 10)
    assert exc.value.http_status == 413


def _scanned_pdf() -> bytes:
    """A one-page PDF whose only content is an image (no text layer)."""
    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 260), False)
    pix.set_rect(pix.irect, (240, 240, 240))
    page.insert_image(page.rect, pixmap=pix)
    return bytes(doc.tobytes())


def test_scanned_pdf_is_unsupported_content_422() -> None:
    with pytest.raises(UnsupportedContent) as exc:
        extract_pdf(_scanned_pdf())
    assert exc.value.http_status == 422
    assert exc.value.code == "unsupported_content"


def test_pdf_page_limit() -> None:
    doc = pymupdf.open()
    for i in range(3):
        doc.new_page().insert_text((72, 72), f"Page {i} with enough text to count as text.")
    with pytest.raises(FileTooLarge):
        extract_pdf(doc.tobytes(), max_pages=2)


def test_corrupt_pdf_is_unsupported_content() -> None:
    with pytest.raises(UnsupportedContent):
        extract_pdf(b"%PDF-1.7\n this is not really a pdf")


# -------------------------------------------------------------------------------------------------
# Extractors on the fixture resumes
# -------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda s: s.id)
def test_extracts_all_key_text(sample: Sample) -> None:
    doc = sample.layout()
    text = doc.text
    gold = sample.gold
    assert gold["personal"]["email"] in text
    assert gold["personal"]["phone"] in text
    # The rendered name may carry a typo; it must still be (nearly) there.
    assert fuzz.partial_ratio(gold["personal"]["name"].lower(), text.lower()) >= 85
    for exp in gold["experience"]:
        assert fuzz.partial_ratio(exp["company"], text) >= 90, exp["company"]


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda s: s.id)
def test_detects_layout(sample: Sample) -> None:
    doc = sample.layout()
    assert doc.layout == EXPECTED_LAYOUT[(doc.file_type, sample.template)]
    assert doc.blocks[0].text.split()[0].lower()[:3] == sample.gold["personal"]["name"].lower()[:3]


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda s: s.id)
def test_reading_order_keeps_each_role_together(sample: Sample) -> None:
    """Each role's heading comes before its own date line, with no other heading in between."""
    lines = [b.text for b in sample.layout().blocks]
    for exp in sample.gold["experience"]:
        heading = next(
            i for i, line in enumerate(lines) if fuzz.partial_ratio(exp["company"], line) >= 90
        )
        assert "–" in lines[heading + 1] or "-" in lines[heading + 1], lines[heading : heading + 2]


def test_pdf_blocks_carry_typography() -> None:
    sample = next(s for s in SAMPLES if s.template == "single_column" and s.path.suffix == ".pdf")
    doc = sample.layout()
    name = doc.blocks[0]
    assert name.bold and name.font_size and name.font_size > (doc.body_font_size or 0)
    assert any(b.kind == "list" for b in doc.blocks)


def test_table_template_reads_row_by_row() -> None:
    """Section titles in the left cell are followed by their own body, not by other titles."""
    sample = next(s for s in SAMPLES if s.template == "table_based")
    lines = [b.text for b in sample.layout().blocks]
    exp_title = next(i for i, ln in enumerate(lines) if ln.endswith("Experience") or "History" in ln)
    assert "|" in lines[exp_title + 1]  # "Title | Company", not another section title


def test_symbol_font_icons_are_dropped() -> None:
    sample = next(s for s in SAMPLES if s.template == "creative" and s.path.suffix == ".pdf")
    assert not any(b.text.startswith("I ") for b in sample.layout().blocks)


def test_docx_text_boxes_and_line_breaks() -> None:
    d = docx.Document()
    d.add_paragraph("Jane Roe")
    p = d.add_paragraph()
    p.add_run("Engineer | Acme").bold = True
    p.add_run("\nJan 2020 – Present")
    host = d.add_paragraph("Before box")
    host._p.append(
        parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox>'
            "<w:txbxContent><w:p><w:r><w:t>Inside box</w:t></w:r></w:p></w:txbxContent>"
            "</v:textbox></v:shape></w:pict></w:r>"
        )
    )
    d.add_heading("Skills", level=1)
    buf = io.BytesIO()
    d.save(buf)
    blocks = extract_docx(buf.getvalue()).blocks
    texts = [b.text for b in blocks]
    assert texts == [
        "Jane Roe",
        "Engineer | Acme",
        "Jan 2020 – Present",
        "Before box",
        "Inside box",
        "Skills",
    ]
    assert blocks[1].bold and not blocks[2].bold
    assert blocks[-1].header_hint


def test_txt_header_hints_and_encodings() -> None:
    raw = "JANE ROE\n========\n--- WORK EXPERIENCE ---\n* Did things\nSkills:\nPython, SQL, AWS\n"
    doc = extract_text(raw.encode("cp1252"))
    hints = {b.text: b.header_hint for b in doc.blocks}
    assert hints == {
        "JANE ROE": True,
        "WORK EXPERIENCE": True,
        "* Did things": False,
        "Skills:": True,
        "Python, SQL, AWS": False,
    }
    assert decode_text("café".encode("utf-16")) == "café"
    assert decode_text("café".encode("cp1252")) == "café"
