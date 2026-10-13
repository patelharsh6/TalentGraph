"""Phase 1.7: render ground-truth profiles into PDF, DOCX and TXT resumes.

    uv run python scripts/render_resumes.py           # render everything, then print stats
    uv run python scripts/render_resumes.py --stats   # stats + exit criterion 1-E (exit 1 on fail)

Templates (6): single_column, two_column, sidebar, table_based, creative (PDF/DOCX) and
minimal_txt (TXT). Each profile is rendered once, with a template and format chosen by a seeded
RNG so that PDF and DOCX each stay >= 40% of the corpus.

Noise (seeded): ~1% character typos, mixed date formats, synonym section headers, 10% of
resumes omit one optional section (also removed from that resume's ground truth), and 30%
have their body sections shuffled.

Output: data/resumes/{train,val,test}/<stem>.<ext> plus <stem>.gt.json, a 70/15/15 split
stratified by (industry, template). The output is byte-for-byte reproducible (fixed PDF
metadata, fixed DOCX zip timestamps). An 18-resume sample (36 files, 3 per template) is copied to
tests/fixtures/resumes/ for unit tests.
"""

from __future__ import annotations

import argparse
import copy
import io
import json
import logging
import random
import shutil
import sys
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.utils import simpleSplit
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from talentgraph.core.schemas.resume import ParsedResume

rl_config.invariant = 1  # fixed creation date and document ID -> reproducible PDF bytes

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger("render_resumes")

ROOT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT_DIR / "data"
PROFILES_DIR = DATA_DIR / "profiles"
RESUMES_DIR = DATA_DIR / "resumes"
FIXTURES_DIR = ROOT_DIR / "tests" / "fixtures" / "resumes"
SPLITS = ("train", "val", "test")
SEED = 42

DOC_TEMPLATES = ["single_column", "two_column", "sidebar", "table_based", "creative"]
TEMPLATES = [*DOC_TEMPLATES, "minimal_txt"]
OPTIONAL_SECTIONS = ["summary", "certifications", "projects"]
BODY_ORDER = ["summary", "experience", "education", "skills", "certifications", "projects"]

HEADER_SYNONYMS = {
    "summary": [
        "Professional Summary",
        "Executive Summary",
        "Career Objective",
        "Profile",
        "About Me",
    ],
    "experience": [
        "Work Experience",
        "Professional Experience",
        "Employment History",
        "Career History",
        "Experience",
    ],
    "education": [
        "Education",
        "Academic Background",
        "Education & Qualifications",
        "Academic Credentials",
    ],
    "skills": [
        "Skills",
        "Technical Skills",
        "Core Competencies",
        "Areas of Expertise",
        "Skills & Proficiencies",
    ],
    "certifications": [
        "Certifications",
        "Certifications & Licenses",
        "Professional Certifications",
        "Credentials",
    ],
    "projects": ["Projects", "Key Projects", "Selected Projects", "Technical Projects"],
}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
DOCX_EPOCH = (1980, 1, 1, 0, 0, 0)


# -------------------------------------------------------------------------------------------------
# Content model shared by every renderer
# -------------------------------------------------------------------------------------------------


@dataclass
class Entry:
    heading: str  # bold line, e.g. "Data Engineer | Acme Corp"
    meta: str = ""  # dates / location line
    bullets: list[str] = field(default_factory=list)


@dataclass
class Section:
    key: str
    title: str
    entries: list[Entry] = field(default_factory=list)
    text: str = ""  # paragraph sections (summary, skills)


@dataclass
class ResumeDoc:
    name: str
    contact: list[str]
    sections: list[Section]


class Noise:
    """All randomness for one resume, drawn from a single seeded RNG."""

    def __init__(self, rng: random.Random, typo_rate: float = 0.01) -> None:
        self.rng = rng
        self.typo_rate = typo_rate

    def typo(self, text: str) -> str:
        chars = list(text)
        for i, ch in enumerate(chars):
            if ch.isalpha() and self.rng.random() < self.typo_rate:
                if i + 1 < len(chars) and chars[i + 1].isalpha():
                    chars[i], chars[i + 1] = chars[i + 1], chars[i]
                else:
                    chars[i] = chr((ord(ch.lower()) - ord("a") + 1) % 26 + ord("a"))
        return "".join(chars)

    def date(self, value: str | None) -> str:
        if not value:
            return "Present"
        year, _, month = value.partition("-")
        if not month.isdigit():
            return value
        m = int(month)
        name = MONTHS[m - 1] if 1 <= m <= 12 else month
        return self.rng.choice(
            [f"{year}-{month}", f"{month}/{year}", f"{name} {year}", f"{name}. {year}", year]
        )

    def header(self, key: str) -> str:
        return self.rng.choice(HEADER_SYNONYMS[key])


def build_document(profile: dict[str, Any], noise: Noise, order: list[str]) -> ResumeDoc:
    personal = profile.get("personal", {})
    loc = personal.get("location", {})
    contact = [
        x
        for x in (
            personal.get("email"),
            personal.get("phone"),
            ", ".join(p for p in (loc.get("city"), loc.get("region")) if p),
            *personal.get("links", []),
        )
        if x
    ]
    built: dict[str, Section] = {}
    if profile.get("summary"):
        built["summary"] = Section(
            "summary", noise.header("summary"), text=noise.typo(profile["summary"])
        )
    if profile.get("experience"):
        sec = Section("experience", noise.header("experience"))
        for exp in profile["experience"]:
            end = "Present" if exp.get("is_current") else noise.date(exp.get("end"))
            sec.entries.append(
                Entry(
                    heading=f"{exp.get('title', '')} | {exp.get('company', '')}",
                    meta=f"{noise.date(exp.get('start'))} – {end} · {exp.get('location', '')}",
                    bullets=[noise.typo(b) for b in exp.get("responsibilities", [])],
                )
            )
        built["experience"] = sec
    if profile.get("education"):
        sec = Section("education", noise.header("education"))
        for edu in profile["education"]:
            meta = f"{noise.date(edu.get('start'))} – {noise.date(edu.get('end'))}"
            if edu.get("grade"):
                meta += f" · {edu['grade']}"
            sec.entries.append(
                Entry(
                    heading=(
                        f"{edu.get('degree', '')} in {edu.get('field', '')}"
                        f" — {edu.get('institution', '')}"
                    ),
                    meta=meta,
                )
            )
        built["education"] = sec
    if profile.get("skills"):
        built["skills"] = Section(
            "skills",
            noise.header("skills"),
            text=", ".join(noise.typo(s.get("raw", "")) for s in profile["skills"]),
        )
    if profile.get("certifications"):
        sec = Section("certifications", noise.header("certifications"))
        for cert in profile["certifications"]:
            meta = f"{cert.get('issuer', '')} · {noise.date(cert.get('date'))}"
            if cert.get("credential_id"):
                meta += f" · ID {cert['credential_id']}"
            sec.entries.append(Entry(heading=cert.get("name", ""), meta=meta))
        built["certifications"] = sec
    if profile.get("projects"):
        sec = Section("projects", noise.header("projects"))
        for prj in profile["projects"]:
            sec.entries.append(
                Entry(
                    heading=prj.get("name", ""),
                    meta=" · ".join(x for x in (prj.get("dates", ""), prj.get("url", "")) if x),
                    bullets=[
                        noise.typo(prj.get("description", "")),
                        f"Technologies: {', '.join(prj.get('technologies', []))}",
                    ],
                )
            )
        built["projects"] = sec
    return ResumeDoc(
        name=noise.typo(personal.get("name", "Candidate")),
        contact=contact,
        sections=[built[k] for k in order if k in built],
    )


# -------------------------------------------------------------------------------------------------
# PDF (ReportLab)
# -------------------------------------------------------------------------------------------------


def _pdf_styles(template: str) -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    accent = colors.HexColor("#2563eb" if template in ("creative", "sidebar") else "#334155")
    return {
        "name": ParagraphStyle(
            "Name",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.white if template == "creative" else colors.HexColor("#0f172a"),
        ),
        "section": ParagraphStyle(
            "Section",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11.5,
            leading=15,
            textColor=accent,
            spaceBefore=7,
            spaceAfter=3,
        ),
        "body": ParagraphStyle(
            "Body", parent=base["Normal"], fontName="Helvetica", fontSize=9, leading=11.5
        ),
        "meta": ParagraphStyle(
            "Meta",
            parent=base["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#64748b"),
        ),
        "side": ParagraphStyle(
            "Side", parent=base["Normal"], fontName="Helvetica", fontSize=8, leading=10.5
        ),
    }


def _section_flowables(sec: Section, st: dict[str, ParagraphStyle], icon: str = "") -> list[Any]:
    out: list[Any] = [Paragraph(escape(f"{icon}{sec.title}"), st["section"])]
    if sec.text:
        out.append(Paragraph(escape(sec.text), st["body"]))
    for e in sec.entries:
        out.append(Paragraph(f"<b>{escape(e.heading)}</b>", st["body"]))
        if e.meta:
            out.append(Paragraph(escape(e.meta), st["meta"]))
        out.extend(Paragraph(escape(f"• {b}"), st["body"]) for b in e.bullets)
        out.append(Spacer(1, 3))
    return out


def render_pdf(doc: ResumeDoc, template: str, out_path: Path) -> None:
    st = _pdf_styles(template)
    pdf = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
        title="Resume",
        author="TalentGraph synthetic corpus",
        creator="TalentGraph",
        producer="TalentGraph",
    )
    width = letter[0] - 72
    story: list[Any] = []
    contact_line = escape(" | ".join(doc.contact))

    if template == "creative":
        band = Table(
            [
                [Paragraph(escape(doc.name), st["name"])],
                [
                    Paragraph(
                        contact_line, ParagraphStyle("c", parent=st["meta"], textColor=colors.white)
                    )
                ],
            ],
            colWidths=[width],
        )
        band.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#1e3a8a")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story += [band, Spacer(1, 8)]
        for sec in doc.sections:
            story += _section_flowables(sec, st, icon="▸ ")
    elif template in ("two_column", "sidebar"):
        side_keys = {"skills", "certifications", "education"}
        side = [s for s in doc.sections if s.key in side_keys]
        main = [s for s in doc.sections if s.key not in side_keys]
        story += [Paragraph(escape(doc.name), st["name"]), Spacer(1, 4)]
        left: list[Any] = [Paragraph("<b>Contact</b>", st["side"])]
        left += [Paragraph(escape(c), st["side"]) for c in doc.contact]
        for sec in side:
            left += _section_flowables(sec, {**st, "body": st["side"], "meta": st["meta"]})
        right: list[Any] = []
        for sec in main:
            right += _section_flowables(sec, st)
        cols = [width * 0.32, width * 0.68] if template == "sidebar" else [width * 0.5, width * 0.5]
        grid = Table([[left, right]], colWidths=cols)
        style = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6)]
        if template == "sidebar":
            style.append(("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#e2e8f0")))
        grid.setStyle(TableStyle(style))
        story.append(grid)
    elif template == "table_based":
        story += [
            Paragraph(escape(doc.name), st["name"]),
            Paragraph(contact_line, st["meta"]),
            Spacer(1, 6),
        ]
        rows = []
        for sec in doc.sections:
            body = _section_flowables(sec, st)[1:]  # title goes in the left cell
            rows.append([Paragraph(f"<b>{escape(sec.title)}</b>", st["body"]), body])
        tbl = Table(rows, colWidths=[width * 0.22, width * 0.78])
        tbl.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(tbl)
    else:  # single_column
        story += [
            Paragraph(escape(doc.name), st["name"]),
            Paragraph(contact_line, st["meta"]),
            HRFlowable(
                width="100%",
                thickness=1,
                color=colors.HexColor("#334155"),
                spaceBefore=4,
                spaceAfter=6,
            ),
        ]
        for sec in doc.sections:
            story += _section_flowables(sec, st)
    pdf.build(story)


# -------------------------------------------------------------------------------------------------
# DOCX (python-docx)
# -------------------------------------------------------------------------------------------------


def _shade(cell: Any, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _docx_section(container: Any, sec: Section, color: RGBColor | None, icon: str = "") -> None:
    title = container.add_paragraph()
    run = title.add_run(f"{icon}{sec.title}")
    run.bold, run.font.size = True, Pt(12)
    if color is not None:
        run.font.color.rgb = color
    if sec.text:
        container.add_paragraph(sec.text)
    for e in sec.entries:
        p = container.add_paragraph()
        p.add_run(e.heading).bold = True
        if e.meta:
            p.add_run(f"\n{e.meta}").italic = True
        for b in e.bullets:
            container.add_paragraph(f"• {b}")


def render_docx(doc: ResumeDoc, template: str, out_path: Path) -> None:
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.5)
    props = d.core_properties
    props.author = props.last_modified_by = "TalentGraph synthetic corpus"
    props.created = props.modified = datetime(2026, 1, 1)
    props.revision = 1

    accent = RGBColor(30, 58, 138) if template in ("creative", "sidebar") else None
    title = d.add_paragraph()
    run = title.add_run(doc.name)
    run.bold, run.font.size = True, Pt(18)
    if template == "creative":
        run.font.color.rgb = RGBColor(30, 58, 138)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact = d.add_paragraph(" | ".join(doc.contact))
    contact.runs[0].italic = True

    if template in ("two_column", "sidebar"):
        side_keys = {"skills", "certifications", "education"}
        table = d.add_table(rows=1, cols=2)
        left, right = table.rows[0].cells
        if template == "sidebar":
            _shade(left, "E2E8F0")
        for sec in doc.sections:
            _docx_section(left if sec.key in side_keys else right, sec, accent)
    elif template == "table_based":
        table = d.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for sec in doc.sections:
            row = table.add_row().cells
            row[0].text = sec.title
            _docx_section(row[1], Section(sec.key, "", sec.entries, sec.text), None)
    else:
        for sec in doc.sections:
            _docx_section(d, sec, accent, icon="▸ " if template == "creative" else "")

    buf = io.BytesIO()
    d.save(buf)
    _write_reproducible_zip(buf.getvalue(), out_path)


def _write_reproducible_zip(data: bytes, out_path: Path) -> None:
    """Re-pack a DOCX with fixed entry timestamps so identical content gives identical bytes."""
    with (
        zipfile.ZipFile(io.BytesIO(data)) as src,
        zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=DOCX_EPOCH)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o644 << 16
            dst.writestr(fixed, src.read(info.filename))


# -------------------------------------------------------------------------------------------------
# TXT
# -------------------------------------------------------------------------------------------------


def render_txt(doc: ResumeDoc, out_path: Path) -> None:
    lines = [doc.name.upper(), "=" * len(doc.name), " | ".join(doc.contact), ""]
    for sec in doc.sections:
        lines.append(f"--- {sec.title.upper()} ---")
        if sec.text:
            lines += (
                simpleSplit(sec.text, "Helvetica", 10, 480) if len(sec.text) > 100 else [sec.text]
            )
        for e in sec.entries:
            lines.append(e.heading)
            if e.meta:
                lines.append(f"  {e.meta}")
            lines += [f"  * {b}" for b in e.bullets]
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8", newline="\n")


# -------------------------------------------------------------------------------------------------
# Corpus
# -------------------------------------------------------------------------------------------------


def assign_template_and_format(rng: random.Random, n: int) -> list[tuple[str, str]]:
    """One (template, format) per profile: 1/6 minimal_txt, the rest alternate PDF/DOCX."""
    plan: list[tuple[str, str]] = []
    doc_formats = ["pdf", "docx"]
    for i in range(n):
        template = TEMPLATES[i % len(TEMPLATES)]
        if template == "minimal_txt":
            plan.append((template, "txt"))
        else:
            plan.append((template, doc_formats[(i // len(TEMPLATES)) % 2]))
    rng.shuffle(plan)
    return plan


def stratified_split(rng: random.Random, keys: dict[str, tuple[str, str]]) -> dict[str, str]:
    """70/15/15 split within every (industry, template) stratum."""
    strata: dict[tuple[str, str], list[str]] = defaultdict(list)
    for stem in sorted(keys):
        strata[keys[stem]].append(stem)
    split_of: dict[str, str] = {}
    for stratum in sorted(strata):
        members = strata[stratum]
        rng.shuffle(members)
        n = len(members)
        n_test = round(n * 0.15)
        n_val = round(n * 0.15)
        for i, stem in enumerate(members):
            split_of[stem] = "test" if i < n_test else "val" if i < n_test + n_val else "train"
    return split_of


def render_all_resumes() -> None:
    profile_files = sorted(PROFILES_DIR.glob("candidate_*.gt.json"))
    if not profile_files:
        raise FileNotFoundError(
            f"No profiles in {PROFILES_DIR}. Run scripts/gen_profiles.py first."
        )

    rng = random.Random(SEED)
    plan = assign_template_and_format(rng, len(profile_files))
    profiles = {
        f.name.removesuffix(".gt.json"): json.loads(f.read_text(encoding="utf-8"))
        for f in profile_files
    }
    stems = sorted(profiles)
    choice = dict(zip(stems, plan, strict=True))
    split_of = stratified_split(
        rng,
        {s: (profiles[s].get("meta", {}).get("industry", "unknown"), choice[s][0]) for s in stems},
    )

    for split in SPLITS:
        out = RESUMES_DIR / split
        if out.exists():
            shutil.rmtree(out)
        out.mkdir(parents=True)

    for stem in stems:
        template, fmt = choice[stem]
        noise = Noise(random.Random(f"{SEED}:{stem}"))  # per-resume RNG: order-independent
        gt = copy.deepcopy(profiles[stem])

        dropped = None
        if noise.rng.random() < 0.10:
            present = [k for k in OPTIONAL_SECTIONS if gt.get(k)]
            if present:
                dropped = noise.rng.choice(present)
                gt[dropped] = None if dropped == "summary" else []
        order = list(BODY_ORDER)
        shuffled = noise.rng.random() < 0.30
        if shuffled:
            noise.rng.shuffle(order)

        doc = build_document(gt, noise, order)
        gt.setdefault("meta", {})
        gt["meta"].update(
            {
                "template": template,
                "format": fmt,
                "split": split_of[stem],
                "noise": {
                    "typo_rate": noise.typo_rate,
                    "dropped_section": dropped,
                    "sections_shuffled": shuffled,
                },
            }
        )
        ParsedResume.model_validate(gt)

        out_dir = RESUMES_DIR / split_of[stem]
        out_doc = out_dir / f"{stem}.{fmt}"
        if fmt == "pdf":
            render_pdf(doc, template, out_doc)
        elif fmt == "docx":
            render_docx(doc, template, out_doc)
        else:
            render_txt(doc, out_doc)
        (out_dir / f"{stem}.gt.json").write_text(
            json.dumps(gt, indent=2) + "\n", encoding="utf-8", newline="\n"
        )

    logger.info("Rendered %d resumes into %s", len(stems), RESUMES_DIR)
    copy_fixture_sample()


def copy_fixture_sample(per_template: int = 3, split: str = "val") -> None:
    """Copy a small sample (all templates, all formats) to tests/fixtures/resumes/."""
    if FIXTURES_DIR.exists():
        shutil.rmtree(FIXTURES_DIR)
    FIXTURES_DIR.mkdir(parents=True)
    taken: Counter[str] = Counter()
    for gt_path in sorted((RESUMES_DIR / split).glob("*.gt.json")):
        meta = json.loads(gt_path.read_text(encoding="utf-8"))["meta"]
        if taken[meta["template"]] >= per_template:
            continue
        stem = gt_path.name.removesuffix(".gt.json")
        shutil.copy2(gt_path, FIXTURES_DIR / gt_path.name)
        shutil.copy2(
            gt_path.with_name(f"{stem}.{meta['format']}"), FIXTURES_DIR / f"{stem}.{meta['format']}"
        )
        taken[meta["template"]] += 1
    logger.info("Copied %d fixture resumes to %s", sum(taken.values()), FIXTURES_DIR)


def print_resume_stats() -> bool:
    """Print corpus statistics and return whether exit criterion 1-E passes."""
    print("\n==================== TALENTGRAPH RESUME CORPUS STATS ====================")
    formats: Counter[str] = Counter()
    templates: Counter[str] = Counter()
    industries: Counter[str] = Counter()
    noise: Counter[str] = Counter()
    total = invalid = missing_gt = 0
    for split in SPLITS:
        s_dir = RESUMES_DIR / split
        docs = (
            sorted(f for f in s_dir.glob("*") if f.suffix in (".pdf", ".docx", ".txt"))
            if s_dir.exists()
            else []
        )
        print(f"Split [{split:5s}]: {len(docs):4d} resumes")
        for d in docs:
            total += 1
            formats[d.suffix[1:]] += 1
            gt_path = d.with_name(f"{d.stem}.gt.json")
            if not gt_path.exists():
                missing_gt += 1
                continue
            try:
                gt = ParsedResume.model_validate_json(gt_path.read_text(encoding="utf-8"))
            except ValueError:
                invalid += 1
                continue
            templates[gt.meta.get("template", "unknown")] += 1
            industries[gt.meta.get("industry", "unknown")] += 1
            n = gt.meta.get("noise", {})
            noise["dropped_section"] += bool(n.get("dropped_section"))
            noise["sections_shuffled"] += bool(n.get("sections_shuffled"))

    def pct(c: int) -> float:
        return c / total * 100 if total else 0.0

    print(
        f"\nTotal resumes: {total}   missing .gt.json: {missing_gt}   invalid .gt.json: {invalid}"
    )
    print(
        "Formats:   "
        + ", ".join(f"{k.upper()} {v} ({pct(v):.1f}%)" for k, v in sorted(formats.items()))
    )
    print("Templates: " + ", ".join(f"{k} {v}" for k, v in sorted(templates.items())))
    print(
        f"Industries ({len(industries)}): "
        + ", ".join(f"{k} {v}" for k, v in sorted(industries.items()))
    )
    print(
        f"Noise: {noise['dropped_section']} dropped a section, "
        f"{noise['sections_shuffled']} shuffled sections"
    )

    passes = (
        total >= 500
        and len(industries) >= 10
        and len(templates) >= 6
        and pct(formats["pdf"]) >= 40
        and pct(formats["docx"]) >= 40
        and missing_gt == 0
        and invalid == 0
    )
    print(f"\nExit criterion 1-E: {'PASS' if passes else 'FAIL'}")
    print("=========================================================================\n")
    return passes


def main() -> int:
    parser = argparse.ArgumentParser(description="Render synthetic resumes into PDF, DOCX and TXT")
    parser.add_argument("--stats", action="store_true", help="Only print corpus statistics")
    args = parser.parse_args()
    if not args.stats:
        render_all_resumes()
    return 0 if print_resume_stats() else 1


if __name__ == "__main__":
    sys.exit(main())
